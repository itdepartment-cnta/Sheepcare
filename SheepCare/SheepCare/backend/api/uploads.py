"""
Uploads API - File upload, parsing, and batch estrus detection.
"""

import os
import uuid
from datetime import datetime
from fastapi import APIRouter, BackgroundTasks, HTTPException, UploadFile, File, Form
from fastapi.responses import FileResponse
from pydantic import BaseModel
from typing import List, Optional
from ..data_access.db_manager import DatabaseManager
from ..data_access.farmcalendar_client import FarmCalendarClient
from ..core.excel_parser import ExcelParser
from ..core.engine import ProcessingEngine
from ..utils.helpers import get_data_dir

router = APIRouter()
db = DatabaseManager()
engine = ProcessingEngine()


class AnimalResultResponse(BaseModel):
    animal_id: str
    result: str
    confidence: float
    probability: float


class UploadResponse(BaseModel):
    upload_id: int
    farm_id: int
    filename: str
    result_file: str = ""
    total_animals: int
    celo_count: int
    no_celo_count: int
    results: List[AnimalResultResponse]


class UploadSummaryResponse(BaseModel):
    id: int
    farm_id: int
    farm_name: str
    filename: str
    upload_date: str
    processed_at: Optional[str]
    total_animals: int
    celo_count: int
    no_celo_count: int


class UploadDetailResponse(UploadSummaryResponse):
    results: List[dict]


def _sync_upload_to_calendar(farm_name: str, celo_count: int, no_celo_count: int,
                              filename: str, animal_details: list) -> None:
    try:
        calendar_client = FarmCalendarClient()
        calendar_client.post_estrus_detection(
            farm_name=farm_name,
            celo_count=celo_count,
            no_celo_count=no_celo_count,
            filename=filename,
            animal_details=animal_details,
        )
    except Exception:
        pass


@router.post("/", response_model=UploadResponse, status_code=201)
async def upload_file(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    farm_id: int = Form(...),
):
    """
    Upload a CSV/Excel file for a specific farm.

    Steps:
    1. Validate farm exists
    2. Parse file (Excel/CSV)
    3. For each animal: save readings, run estrus detection
    4. Save detection results
    5. Return summary
    """
    # 1. Validate farm
    farm = db.get_farm(farm_id)
    if not farm:
        raise HTTPException(status_code=404, detail="Farm not found")

    # Ensure farm is synced to Farm Calendar (deferred — does not block upload)
    if not db.get_fc_uuid("farm", farm_id):
        from ..api.farms import _sync_farm_to_calendar
        background_tasks.add_task(_sync_farm_to_calendar, farm_id,
                                  farm["name"], farm.get("location", ""))

    # 2. Validate file type
    if not file.filename or not file.filename.lower().endswith(
        (".xlsx", ".xls", ".csv")
    ):
        raise HTTPException(
            status_code=400,
            detail="Invalid file format. Only .xlsx, .xls, and .csv are supported.",
        )

    # 3. Save file temporarily
    data_dir = get_data_dir()
    uploads_dir = os.path.join(data_dir, "uploads")
    os.makedirs(uploads_dir, exist_ok=True)

    safe_filename = f"{uuid.uuid4().hex}_{file.filename}"
    file_path = os.path.join(uploads_dir, safe_filename)

    content = await file.read()
    with open(file_path, "wb") as f:
        f.write(content)

    # 4. Parse file
    try:
        animals_data = ExcelParser.parse_content(content, file.filename)
    except Exception as e:
        # Clean up on error
        if os.path.exists(file_path):
            os.remove(file_path)
        raise HTTPException(
            status_code=422,
            detail=f"Failed to parse file: {str(e)}",
        )

    if not animals_data:
        if os.path.exists(file_path):
            os.remove(file_path)
        raise HTTPException(
            status_code=422,
            detail="No valid animal data found in file. Check that column A has animal IDs and columns C+ have numeric values.",
        )

    # 5. Create upload record
    upload_id = db.create_upload(farm_id, file.filename)

    # 6. Process each animal
    results = []
    celo_count = 0
    no_celo_count = 0

    for animal_data in animals_data:
        animal_id = animal_data["animal_id"]
        readings = animal_data["readings"]

        if not readings:
            continue

        # Register or find animal
        db_animal_id = db.add_animal(farm_id, animal_id, name=animal_id)

        # Animal FC sync is handled in background by _sync_upload_to_calendar

        # Save readings
        for reading in readings:
            db.save_animal_reading(
                db_animal_id,
                reading["date"],
                reading["value"],
                source_file=file.filename,
            )

        # Run detection on the time series
        values = [r["value"] for r in readings]

        # Pasar valores crudos directamente al modelo (sin normalizar).
        # El modelo fue entrenado con las lecturas RAW del Excel, no con
        # señales normalizadas. Solo limpiamos NaN/inf.
        import numpy as np

        raw_array = np.array(values, dtype=np.float64)
        raw_array = raw_array[np.isfinite(raw_array)]

        # El modelo espera exactamente 5 lecturas por animal
        if len(raw_array) != 5:
            continue

        try:
            detection = engine.predict_estrus(raw_array)
        except Exception:
            # Si falla la predicción, ignoramos este animal
            continue

        result_label = "Celo" if detection["estrus_detected"] else "No celo"
        if detection["estrus_detected"]:
            celo_count += 1
        else:
            no_celo_count += 1

        # Save result
        db.save_detection_result(
            db_animal_id,
            upload_id,
            result_label,
            detection["confidence"],
            detection["probability"],
        )

        results.append(
            AnimalResultResponse(
                animal_id=animal_id,
                result=result_label,
                confidence=round(detection["confidence"], 4),
                probability=round(detection["probability"], 4),
            )
        )

    # 7. Finalize upload record
    db.finalize_upload(upload_id, len(animals_data), celo_count, no_celo_count)

    # 8. Generate result Excel
    result_file_name = ""
    if file.filename.lower().endswith((".xlsx", ".xls")):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment

            result_filename = f"{uuid.uuid4().hex}_resultado_{file.filename}"
            result_path = os.path.join(uploads_dir, result_filename)

            wb = openpyxl.load_workbook(file_path)
            ws = wb.active if wb.active else wb[wb.sheetnames[0]]

            # Find the last column with data
            max_col = ws.max_column or 1
            result_col = max_col + 1

            # Add header "Resultado"
            header_cell = ws.cell(row=1, column=result_col)
            header_cell.value = "Resultado"
            header_cell.font = Font(bold=True, color="FFFFFF")
            header_cell.fill = PatternFill(
                start_color="E28474", end_color="E28474", fill_type="solid"
            )
            header_cell.alignment = Alignment(horizontal="center")

            # Build lookup: animal_id -> result
            result_map = {r.animal_id: r.result for r in results}

            # Fill results for each row
            for row_idx in range(2, ws.max_row + 1):
                animal_id = str(ws.cell(row=row_idx, column=1).value or "").strip()
                result_val = result_map.get(animal_id, "")
                cell = ws.cell(row=row_idx, column=result_col)
                cell.value = result_val
                cell.alignment = Alignment(horizontal="center")
                if result_val == "Celo":
                    cell.font = Font(color="C0503A", bold=True)
                    cell.fill = PatternFill(
                        start_color="FDE8E8", end_color="FDE8E8", fill_type="solid"
                    )
                elif result_val == "No celo":
                    cell.font = Font(color="6B7280")
                    cell.fill = PatternFill(
                        start_color="F3F4F6", end_color="F3F4F6", fill_type="solid"
                    )

            wb.save(result_path)
            result_file_name = result_filename
        except Exception:
            pass  # Non-critical: continue even if Excel generation fails

    # 9. Clean up temp file
    if os.path.exists(file_path):
        os.remove(file_path)

    # 10. Post to Farm Calendar in background (non-blocking)
    background_tasks.add_task(
        _sync_upload_to_calendar,
        farm_name=farm.get("name", f"Farm {farm_id}"),
        celo_count=celo_count,
        no_celo_count=no_celo_count,
        filename=file.filename,
        animal_details=[
            {"animal_id": r.animal_id, "result": r.result, "confidence": r.confidence}
            for r in results
        ],
    )

    return UploadResponse(
        upload_id=upload_id,
        farm_id=farm_id,
        filename=file.filename,
        result_file=result_file_name,
        total_animals=len(animals_data),
        celo_count=celo_count,
        no_celo_count=no_celo_count,
        results=results,
    )


@router.get("/", response_model=List[UploadSummaryResponse])
async def list_uploads(farm_id: Optional[int] = None):
    """List all uploads, optionally filtered by farm."""
    uploads = db.list_uploads(farm_id)
    return [UploadSummaryResponse(**u) for u in uploads]


@router.get("/{upload_id}", response_model=UploadDetailResponse)
async def get_upload(upload_id: int):
    """Get upload details with detection results."""
    upload = db.get_upload(upload_id)
    if not upload:
        raise HTTPException(status_code=404, detail="Upload not found")

    results = db.get_results(upload_id=upload_id)
    return UploadDetailResponse(
        **upload,
        results=[dict(r) for r in results],
    )


@router.get("/download/{filename}")
async def download_result(filename: str):
    """Download a generated result file."""
    data_dir = get_data_dir()
    uploads_dir = os.path.join(data_dir, "uploads")
    file_path = os.path.join(uploads_dir, filename)
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(file_path, filename=filename)


@router.get("/{upload_id}/results", response_model=List[AnimalResultResponse])
async def get_upload_results(upload_id: int):
    """Get detection results for a specific upload."""
    upload = db.get_upload(upload_id)
    if not upload:
        raise HTTPException(status_code=404, detail="Upload not found")

    results = db.get_results(upload_id=upload_id)
    return [
        AnimalResultResponse(
            animal_id=r["animal_tag"],
            result=r["result"],
            confidence=r["confidence"],
            probability=r["probability"],
        )
        for r in results
    ]
