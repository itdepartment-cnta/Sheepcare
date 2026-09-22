"""
Utility helpers for SHEEPCARE backend.
"""

import os
import sys
from datetime import datetime, timedelta
from typing import Dict, List, Optional, Tuple


def get_project_root() -> str:
    """Return the absolute path to the project root directory."""
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def get_data_dir() -> str:
    """Return the path to the data directory, creating it if needed."""
    db_path = os.environ.get("SHEEPCARE_DB_PATH")
    if db_path:
        data_dir = os.path.dirname(db_path)
    else:
        data_dir = os.path.join(get_project_root(), "data")
    os.makedirs(data_dir, exist_ok=True)
    return data_dir


def resource_path(relative_path: str) -> str:
    """
    Get absolute path to a resource, works for dev and for PyInstaller.
    """
    try:
        base_path = sys._MEIPASS  # type: ignore[attr-defined]
    except AttributeError:
        base_path = get_project_root()
    return os.path.join(base_path, relative_path)


def format_timestamp(dt_str: str, fmt: str = "%Y-%m-%d %H:%M") -> str:
    """Format an ISO datetime string for display."""
    try:
        dt = datetime.fromisoformat(dt_str)
        return dt.strftime(fmt)
    except (ValueError, TypeError):
        return dt_str


def format_confidence(confidence: float) -> str:
    """Format confidence score as percentage string."""
    return f"{confidence * 100:.1f}%"


def append_result_column(
    source_path: str,
    dest_path: str,
    id_to_result: Dict[str, str],
    result_colors: Dict[str, Tuple[str, str]],
) -> None:
    """
    Copy an uploaded workbook and append a "Resultado" column, filled by
    looking up each row's column-A value in `id_to_result`.

    `result_colors` maps a result label to (font_color, fill_color) hex pairs
    (without '#'); a label not present there is left with default styling.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.load_workbook(source_path)
    ws = wb.active if wb.active else wb[wb.sheetnames[0]]

    max_col = ws.max_column or 1
    result_col = max_col + 1

    header_cell = ws.cell(row=1, column=result_col)
    header_cell.value = "Resultado"
    header_cell.font = Font(bold=True, color="FFFFFF")
    header_cell.fill = PatternFill(start_color="E28474", end_color="E28474", fill_type="solid")
    header_cell.alignment = Alignment(horizontal="center")

    for row_idx in range(2, ws.max_row + 1):
        key = str(ws.cell(row=row_idx, column=1).value or "").strip()
        result_val = id_to_result.get(key, "")
        cell = ws.cell(row=row_idx, column=result_col)
        cell.value = result_val
        cell.alignment = Alignment(horizontal="center")
        colors = result_colors.get(result_val)
        if colors:
            font_color, fill_color = colors
            cell.font = Font(color=font_color, bold=True)
            cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

    wb.save(dest_path)


def date_range_filter(
    readings: List[dict],
    days: Optional[int] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
) -> List[dict]:
    """
    Filter readings list by date range.
    If `days` is provided, filter to last N days.
    """
    if not readings:
        return []

    if days is not None:
        cutoff = datetime.now() - timedelta(days=days)
        return [
            r
            for r in readings
            if datetime.fromisoformat(r.get("timestamp", "")).replace(tzinfo=None)
            >= cutoff
        ]

    if start_date:
        start = datetime.fromisoformat(start_date)
        readings = [
            r
            for r in readings
            if datetime.fromisoformat(r.get("timestamp", "")).replace(tzinfo=None)
            >= start
        ]

    if end_date:
        end = datetime.fromisoformat(end_date)
        readings = [
            r
            for r in readings
            if datetime.fromisoformat(r.get("timestamp", "")).replace(tzinfo=None)
            <= end
        ]

    return readings
